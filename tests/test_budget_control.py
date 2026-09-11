"""Test Budget Control — Multi-Level Approval + FA Acknowledge + Level Management +
Excel Upload (Phase 10, 2026-09-09, Business Decision v4.1)

ครอบคลุม Logic ที่เสี่ยงที่สุดตาม Design: การเรียง Level (ข้ามลำดับไม่ได้), จุดหักยอด
งบมีจุดเดียว (FA Acknowledge เท่านั้น), Revise คืนยอด+Reset กลับ Level 1 ใหม่เสมอ, สิทธิ์
(ผู้อนุมัติเจาะจงคน + Admin Override), แผนกไม่มี Level ข้ามตรงไป FA, เกินงบ (Block/Force)
"""

from __future__ import annotations

from datetime import date
from decimal import Decimal

import openpyxl
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.core.security import hash_password
from app.models import ApprovalRequest, ARBudgetType, BudgetApprovalLevel, BudgetMaster, User
from app.services.ar_pdf import render_ar_html
from tests.test_approval_requests import _sample_ar_body

_PASSWORD = "password123456"


def _make_user(db_session: Session, **kwargs) -> User:
    kwargs.setdefault("password_hash", hash_password(_PASSWORD))
    # Full RBAC (Correction 2026-09-10): Default can_view_pr/can_view_ar = True ให้ Fixture
    # นี้ (เหมือน conftest.plain_user — ผู้ใช้ทั่วไปในระบบมีสิทธิ์ PR/AR ของตัวเองอยู่แล้ว)
    # Test ที่ต้องการยืนยัน Permission Gate เอง (403 กรณีไม่มีสิทธิ์) Override ผ่าน kwargs
    kwargs.setdefault("can_view_pr", True)
    kwargs.setdefault("can_view_ar", True)
    user = User(**kwargs)
    db_session.add(user)
    db_session.commit()
    db_session.refresh(user)
    return user


def _login_as(client: TestClient, email: str, password: str = _PASSWORD) -> None:
    res = client.post("/auth/login", json={"email": email, "password": password})
    assert res.status_code == 200, res.text


def _make_budget_master(db_session: Session, **kwargs) -> BudgetMaster:
    # budget_no คือ Key จริงที่ไม่ซ้ำกัน (Correction 2026-09-09) — Default ตรงกับ
    # budget_no Default ของ _sample_ar_body() ใน test_approval_requests.py ("5100-01")
    # เพื่อให้ Match กันอัตโนมัติในเทสที่ไม่ได้ Override ค่านี้ — account_code เป็นแค่
    # รหัสบัญชี/หมวดหมู่ ไม่ใช่ Key ซ้ำกันได้ตามจริง
    defaults = dict(
        budget_no="5100-01",
        department="Production",
        budget_type="expenses",
        account_code="5100",
        period_start=date(2026, 1, 1),
        period_end=date(2026, 12, 31),
        budgeted_amount=Decimal("10000.00"),
        used_amount=Decimal("0.00"),
    )
    defaults.update(kwargs)
    defaults["budget_type"] = ARBudgetType(defaults["budget_type"])
    row = BudgetMaster(**defaults)
    db_session.add(row)
    db_session.commit()
    db_session.refresh(row)
    return row


def _make_level(db_session: Session, **kwargs) -> BudgetApprovalLevel:
    row = BudgetApprovalLevel(**kwargs)
    db_session.add(row)
    db_session.commit()
    db_session.refresh(row)
    return row


def _submit_for_approval(client: TestClient, ar_id: int) -> dict:
    """Correction 2026-09-10: เดิมพิมพ์/ดาวน์โหลด PDF ครั้งแรกคือจุด Finalize+เริ่ม
    Workflow (ตั้งชื่อ Helper เดิมว่า _finalize) เปลี่ยนมาเป็น Endpoint แยก
    submit-for-approval แล้ว — AR อาจยังเป็น Draft อยู่หลังเรียก Helper นี้ก็ได้ถ้า
    แผนกมี Level ให้รอ (จะ Finalize จริงก็ต่อเมื่อ Level แรกอนุมัติผ่าน หรือแผนกไม่มี
    Level เลย) ดู Docstring submit_ar_for_approval ใน approval_requests.py"""
    res = client.post(f"/ars/{ar_id}/submit-for-approval")
    assert res.status_code == 200, res.text
    return client.get(f"/ars/{ar_id}").json()


# ───────────────────────── Level Management CRUD ─────────────────────────
def test_level_management_requires_admin(client: TestClient, plain_user: User):
    _login_as(client, plain_user.email, "plainpass123")
    res = client.post(
        "/budget-approval-levels",
        json={
            "department": "Production",
            "level_no": 1,
            "level_name": "Manager",
            "approver_user_id": plain_user.id,
        },
    )
    assert res.status_code == 403


def test_level_management_crud(client: TestClient, admin_user: User, db_session: Session):
    manager = _make_user(
        db_session, name="Manager A", email="mgr@example.com", department="Production"
    )
    _login_as(client, admin_user.email, "adminpass123")

    res = client.post(
        "/budget-approval-levels",
        json={
            "department": "Production",
            "level_no": 1,
            "level_name": "Manager",
            "approver_user_id": manager.id,
        },
    )
    assert res.status_code == 201, res.text
    level = res.json()
    assert level["approver_name"] == "Manager A"

    # ห้ามคนเดิมซ้ำใน Level เดียวกัน (Multi-approver per Level, OR — Correction
    # 2026-09-09) แม้ level_name ที่ส่งมาจะไม่ตรงกับที่ตั้งไว้ก็ยัง 409 เพราะ Reason
    # แรกที่เจอคือ "คนซ้ำ" (ดู _validate_level_slot ใน budget_levels.py)
    dup = client.post(
        "/budget-approval-levels",
        json={
            "department": "Production",
            "level_no": 1,
            "level_name": "Manager (dup)",
            "approver_user_id": manager.id,
        },
    )
    assert dup.status_code == 409

    patched = client.patch(
        f"/budget-approval-levels/{level['id']}", json={"level_name": "Manager Approve"}
    )
    assert patched.status_code == 200
    assert patched.json()["level_name"] == "Manager Approve"

    deleted = client.delete(f"/budget-approval-levels/{level['id']}")
    assert deleted.status_code == 204
    listed = client.get("/budget-approval-levels", params={"department": "Production"}).json()
    assert listed[0]["is_active"] is False


# ───────────────── /departments แสดงแผนกทั้งหมด ไม่ใช่แค่แผนกที่ตั้ง Level ไว้แล้ว ─────────────────
# (2026-09-10) แก้ตามคำขอ: หน้า จัดการ Level ต้องโชว์แผนกทั้งหมดให้เลือกทันที ไม่ต้องพิมพ์
# เดา — รวมแผนกที่ "มีแค่ User" แต่ยังไม่เคยตั้ง Level เลยด้วย ไม่ใช่รวมแค่แผนกที่มี Level
def test_departments_endpoint_includes_depts_with_users_but_no_levels(
    client: TestClient, admin_user: User, db_session: Session
):
    _make_user(db_session, name="Only User", email="onlyuser@example.com", department="Warehouse")
    manager = _make_user(
        db_session, name="Manager Dept", email="mgrdept@example.com", department="Production"
    )
    _login_as(client, admin_user.email, "adminpass123")
    client.post(
        "/budget-approval-levels",
        json={
            "department": "Production",
            "level_no": 1,
            "level_name": "Manager",
            "approver_user_id": manager.id,
        },
    )

    res = client.get("/budget-approval-levels/departments")
    assert res.status_code == 200
    depts = res.json()
    assert "Production" in depts  # มี Level ตั้งไว้แล้ว
    assert "Warehouse" in depts  # ยังไม่มี Level เลย แต่มี User สังกัดแผนกนี้ ต้องโชว์ด้วย


# ───────────────── Multi-approver per Level (OR) — Correction 2026-09-09 ─────────────────
def test_level_management_multiple_approvers_same_level(
    client: TestClient, admin_user: User, db_session: Session
):
    """Level เดียวกัน (department+level_no) ตั้งผู้อนุมัติได้มากกว่า 1 คน ตราบใดที่
    level_name ตรงกันเป๊ะ — ตรงกับตัวอย่าง Approve Flow จริงที่ผู้ใช้ส่งมา (Level
    "President or Director" มี Hori/Ochi/Ukai พร้อมกัน)"""
    hori = _make_user(db_session, name="Hori", email="hori@example.com")
    ochi = _make_user(db_session, name="Ochi", email="ochi@example.com")
    _login_as(client, admin_user.email, "adminpass123")

    first = client.post(
        "/budget-approval-levels",
        json={
            "department": "Production",
            "level_no": 4,
            "level_name": "President or Director",
            "approver_user_id": hori.id,
        },
    )
    assert first.status_code == 201, first.text

    # เพิ่มคนที่ 2 เข้า Level เดียวกัน (level_name ตรงกันเป๊ะ) — ต้องผ่าน ไม่ใช่ 409 แบบเดิม
    second = client.post(
        "/budget-approval-levels",
        json={
            "department": "Production",
            "level_no": 4,
            "level_name": "President or Director",
            "approver_user_id": ochi.id,
        },
    )
    assert second.status_code == 201, second.text

    listed = client.get("/budget-approval-levels", params={"department": "Production"}).json()
    level4_rows = [lv for lv in listed if lv["level_no"] == 4]
    assert len(level4_rows) == 2
    assert {lv["approver_name"] for lv in level4_rows} == {"Hori", "Ochi"}


def test_level_management_mismatched_level_name_in_same_level_rejected(
    client: TestClient, admin_user: User, db_session: Session
):
    """คนที่ 2 ใน Level เดียวกันต้องตั้งชื่อ Level ให้ตรงกับกลุ่มเดิมเป๊ะ — กัน Audit
    Trail สับสนว่า Level ไหนชื่ออะไรกันแน่"""
    hori = _make_user(db_session, name="Hori2", email="hori2@example.com")
    ochi = _make_user(db_session, name="Ochi2", email="ochi2@example.com")
    _login_as(client, admin_user.email, "adminpass123")

    client.post(
        "/budget-approval-levels",
        json={
            "department": "Sales",
            "level_no": 4,
            "level_name": "President or Director",
            "approver_user_id": hori.id,
        },
    )
    mismatched = client.post(
        "/budget-approval-levels",
        json={
            "department": "Sales",
            "level_no": 4,
            "level_name": "President/Director",  # สะกดไม่ตรงกับกลุ่มเดิมเป๊ะ
            "approver_user_id": ochi.id,
        },
    )
    assert mismatched.status_code == 409


def test_level_management_rename_group_syncs_all_rows(
    client: TestClient, admin_user: User, db_session: Session
):
    """แก้ level_name ผ่านแถวใดแถวหนึ่งในกลุ่ม ต้อง Sync ชื่อใหม่ให้ทุกคนในกลุ่มเดียวกัน
    อัตโนมัติ (department+level_no เดียวกัน)"""
    hori = _make_user(db_session, name="Hori3", email="hori3@example.com")
    ochi = _make_user(db_session, name="Ochi3", email="ochi3@example.com")
    _login_as(client, admin_user.email, "adminpass123")

    r1 = client.post(
        "/budget-approval-levels",
        json={
            "department": "Maintenance",
            "level_no": 4,
            "level_name": "President or Director",
            "approver_user_id": hori.id,
        },
    ).json()
    client.post(
        "/budget-approval-levels",
        json={
            "department": "Maintenance",
            "level_no": 4,
            "level_name": "President or Director",
            "approver_user_id": ochi.id,
        },
    )

    renamed = client.patch(
        f"/budget-approval-levels/{r1['id']}", json={"level_name": "President / Director"}
    )
    assert renamed.status_code == 200

    listed = client.get("/budget-approval-levels", params={"department": "Maintenance"}).json()
    assert all(lv["level_name"] == "President / Director" for lv in listed)


def test_multi_approver_level_any_one_can_approve(
    client: TestClient, plain_user: User, db_session: Session
):
    """Level มี 2 คน (OR) — ใครกดอนุมัติก่อนก็ผ่านได้ ไม่ต้องรอครบทุกคน และคนที่เหลือใน
    กลุ่มกดซ้ำไม่ได้อีกต่อไปเพราะ Level เดินหน้าไปแล้ว"""
    manager_a = _make_user(
        db_session, name="Manager A2", email="mgra2@example.com", department="Production"
    )
    manager_b = _make_user(
        db_session, name="Manager B2", email="mgrb2@example.com", department="Production"
    )
    _make_user(db_session, name="FA B2", email="fab2@example.com", is_fa=True)
    _make_level(
        db_session,
        department="Production",
        level_no=1,
        level_name="Manager",
        approver_user_id=manager_a.id,
    )
    _make_level(
        db_session,
        department="Production",
        level_no=1,
        level_name="Manager",
        approver_user_id=manager_b.id,
    )
    _make_budget_master(db_session)

    _login_as(client, plain_user.email, "plainpass123")
    ar_id = client.post("/ars", json=_sample_ar_body()).json()["id"]
    _submit_for_approval(client, ar_id)

    # ก่อนมีใครกด — Progress ต้องโชว์ทั้งคู่เป็นผู้มีสิทธิ์ Level นี้
    progress_before = client.get(f"/ars/{ar_id}/approval-progress").json()
    level1_step = next(s for s in progress_before if s["level_no"] == 1)
    assert set(level1_step["approver_user_ids"]) == {manager_a.id, manager_b.id}
    assert "Manager A2" in level1_step["approver_names"]
    assert "Manager B2" in level1_step["approver_names"]

    # manager_b (ไม่ใช่คนแรกที่ตั้งไว้) กดอนุมัติก่อน — ต้องผ่านได้เลย (OR)
    _login_as(client, manager_b.email)
    approved = client.post(f"/ars/{ar_id}/approve-level")
    assert approved.status_code == 200, approved.text
    assert approved.json()["budget_approval_status"] == "pending_fa_acknowledge"

    # manager_a (คนที่เหลือในกลุ่มเดียวกัน) กดซ้ำไม่ได้อีกแล้ว เพราะ Workflow เดินหน้า
    # ไปพ้น Level 1 แล้ว (ไม่ใช่ current_approval_level อีกต่อไป)
    _login_as(client, manager_a.email)
    stale = client.post(f"/ars/{ar_id}/approve-level")
    assert stale.status_code == 409

    progress_after = client.get(f"/ars/{ar_id}/approval-progress").json()
    level1_after = next(s for s in progress_after if s["level_no"] == 1)
    assert level1_after["status"] == "approved"
    assert level1_after["acted_by_name"] == "Manager B2"


# ───────────────────────── Requester ต้องมี Department ─────────────────────────
def test_create_ar_requires_department(client: TestClient, db_session: Session):
    no_dept_user = _make_user(
        db_session, name="No Dept", email="nodept@example.com", department=None
    )
    _login_as(client, no_dept_user.email)
    res = client.post("/ars", json=_sample_ar_body())
    assert res.status_code == 400


# ───────────────────────── Multi-Level -> FA Acknowledge (Happy Path) ─────────────────────────
def test_two_level_then_fa_acknowledge_deducts_budget_once(
    client: TestClient, plain_user: User, db_session: Session
):
    manager = _make_user(
        db_session, name="Manager A", email="mgr2@example.com", department="Production"
    )
    gm = _make_user(db_session, name="GM A", email="gm2@example.com", department="Production")
    fa = _make_user(db_session, name="FA A", email="fa2@example.com", is_fa=True)
    _make_level(
        db_session,
        department="Production",
        level_no=1,
        level_name="Manager",
        approver_user_id=manager.id,
    )
    _make_level(
        db_session,
        department="Production",
        level_no=2,
        level_name="General Manager",
        approver_user_id=gm.id,
    )
    master = _make_budget_master(db_session)

    _login_as(client, plain_user.email, "plainpass123")
    created = client.post("/ars", json=_sample_ar_body()).json()
    ar_id = created["id"]

    ar_after_finalize = _submit_for_approval(client, ar_id)
    assert ar_after_finalize["budget_approval_status"] == "pending"
    assert ar_after_finalize["current_approval_level"] == 1
    assert ar_after_finalize["budget_master_id"] == master.id
    # Feedback ผู้ใช้ 2026-09-11: Badge "รออนุมัติ Level" เดิมไม่บอกว่ารอ Level ไหน — GET
    # /ars/{id} ต้อง Resolve ชื่อ Level จริงมาให้ (current_level_name) เหมือนที่
    # MyApprovalItem มีอยู่แล้ว ไม่ใช่แค่เลข current_approval_level เฉยๆ
    assert ar_after_finalize["current_level_name"] == "Manager"

    # plain_user (ไม่ใช่ผู้อนุมัติ Level 1) กด Approve ไม่ได้
    forbidden = client.post(f"/ars/{ar_id}/approve-level")
    assert forbidden.status_code == 403

    _login_as(client, manager.email)
    step1 = client.post(f"/ars/{ar_id}/approve-level")
    assert step1.status_code == 200, step1.text
    assert step1.json()["current_approval_level"] == 2
    assert step1.json()["budget_approval_status"] == "pending"
    assert step1.json()["current_level_name"] == "General Manager"

    # ข้าม Level ไม่ได้ — gm ต้องรอจนกว่าจะถึงคิวตัวเอง (แต่ในเคสนี้ถึงคิวแล้วพอดี) ลอง
    # ให้ Manager กด Approve ซ้ำ (ไม่ใช่ผู้อนุมัติ Level ปัจจุบันอีกแล้ว) ต้องโดนบล็อก
    _login_as(client, manager.email)
    stale = client.post(f"/ars/{ar_id}/approve-level")
    assert stale.status_code == 403

    _login_as(client, gm.email)
    step2 = client.post(f"/ars/{ar_id}/approve-level")
    assert step2.status_code == 200, step2.text
    assert step2.json()["budget_approval_status"] == "pending_fa_acknowledge"
    assert step2.json()["current_approval_level"] is None
    assert step2.json()["current_level_name"] == "FA Acknowledge"

    _login_as(client, plain_user.email, "plainpass123")
    forbidden_fa = client.post(f"/ars/{ar_id}/fa-acknowledge", json={"force": False})
    assert forbidden_fa.status_code == 403

    _login_as(client, fa.email)
    ack = client.post(f"/ars/{ar_id}/fa-acknowledge", json={"force": False})
    assert ack.status_code == 200, ack.text
    body = ack.json()
    assert body["budget_approval_status"] == "approved"
    assert Decimal(str(body["budget_deducted_amount"])) == Decimal("9327.10")
    assert body["budget_overridden"] is False

    db_session.refresh(master)
    assert master.used_amount == Decimal("9327.10")

    progress = client.get(f"/ars/{ar_id}/approval-progress").json()
    assert [s["status"] for s in progress] == ["approved", "approved", "approved"]
    assert progress[-1]["step_type"] == "fa_acknowledge"

    # PDF Auto-fill: ชื่อผู้อนุมัติต้องปรากฏใน HTML ที่ Render จริง (ตาราง Authority + F&A)
    ar_obj = db_session.get(ApprovalRequest, ar_id)
    html = render_ar_html(db_session, ar_obj)
    assert "Manager A" in html
    assert "GM A" in html
    assert "FA A" in html


# ───────────────────────── Admin Override ─────────────────────────
def test_admin_can_override_level_approval(
    client: TestClient, plain_user: User, admin_user: User, db_session: Session
):
    manager = _make_user(
        db_session, name="Manager B", email="mgr3@example.com", department="Production"
    )
    _make_level(
        db_session,
        department="Production",
        level_no=1,
        level_name="Manager",
        approver_user_id=manager.id,
    )
    _make_budget_master(db_session)

    _login_as(client, plain_user.email, "plainpass123")
    ar_id = client.post("/ars", json=_sample_ar_body()).json()["id"]
    _submit_for_approval(client, ar_id)

    _login_as(client, admin_user.email, "adminpass123")
    res = client.post(f"/ars/{ar_id}/approve-level")
    assert res.status_code == 200

    progress = client.get(f"/ars/{ar_id}/approval-progress").json()
    assert progress[0]["acted_as_override"] is True
    assert progress[0]["acted_by_name"] == admin_user.name


# ───────────────────────── Reject -> Revise -> เข้าคิว Level 1 ใหม่ ─────────────────────────
def test_reject_then_revise_resets_to_level_one(
    client: TestClient, plain_user: User, db_session: Session
):
    manager = _make_user(
        db_session, name="Manager C", email="mgr4@example.com", department="Production"
    )
    gm = _make_user(db_session, name="GM C", email="gm4@example.com", department="Production")
    _make_level(
        db_session,
        department="Production",
        level_no=1,
        level_name="Manager",
        approver_user_id=manager.id,
    )
    _make_level(
        db_session,
        department="Production",
        level_no=2,
        level_name="General Manager",
        approver_user_id=gm.id,
    )
    _make_budget_master(db_session)

    _login_as(client, plain_user.email, "plainpass123")
    ar_id = client.post("/ars", json=_sample_ar_body()).json()["id"]
    _submit_for_approval(client, ar_id)

    _login_as(client, manager.email)
    client.post(f"/ars/{ar_id}/approve-level")

    _login_as(client, gm.email)
    rejected = client.post(f"/ars/{ar_id}/reject-level", json={"reason": "ยอดไม่ตรงกับใบเสนอราคา"})
    assert rejected.status_code == 200
    assert rejected.json()["budget_approval_status"] == "rejected"

    _login_as(client, plain_user.email, "plainpass123")
    revised = client.post(f"/ars/{ar_id}/revise")
    assert revised.status_code == 201
    revised_body = revised.json()
    assert revised_body["budget_approval_status"] == "not_submitted"
    assert revised_body["current_approval_level"] is None

    # Finalize ฉบับ Revise แล้ว — ต้องเริ่มที่ Level 1 ใหม่ (ไม่ข้ามไป Level 2 เพราะ
    # Level 1 เคยผ่านมาก่อนหน้า Reject)
    ar_after = _submit_for_approval(client, revised_body["id"])
    assert ar_after["budget_approval_status"] == "pending"
    assert ar_after["current_approval_level"] == 1


# ───────────────────────── แผนกไม่มี Level -> ข้ามตรงไป FA ─────────────────────────
def test_department_with_no_levels_skips_straight_to_fa(client: TestClient, db_session: Session):
    fa = _make_user(db_session, name="FA Sales", email="fasales@example.com", is_fa=True)
    requester = _make_user(
        db_session, name="Sales Req", email="salesreq@example.com", department="Sales"
    )
    _make_budget_master(db_session, department="Sales", budget_no="6100-01", account_code="6100")

    _login_as(client, requester.email)
    ar_id = client.post("/ars", json=_sample_ar_body(budget_no="6100-01")).json()["id"]
    ar_after = _submit_for_approval(client, ar_id)
    assert ar_after["budget_approval_status"] == "pending_fa_acknowledge"
    assert ar_after["current_approval_level"] is None

    _login_as(client, fa.email)
    ack = client.post(f"/ars/{ar_id}/fa-acknowledge", json={"force": False})
    assert ack.status_code == 200
    assert ack.json()["budget_approval_status"] == "approved"


# ───────────────────────── เกินงบ: Block ไม่ Force / Force สำเร็จ ─────────────────────────
def test_fa_acknowledge_over_budget_blocked_unless_forced(client: TestClient, db_session: Session):
    fa = _make_user(db_session, name="FA Tight", email="fatight@example.com", is_fa=True)
    requester = _make_user(
        db_session, name="Tight Req", email="tightreq@example.com", department="Tight"
    )
    _make_budget_master(
        db_session,
        department="Tight",
        budget_no="7100-01",
        account_code="7100",
        budgeted_amount=Decimal("1000.00"),
    )

    _login_as(client, requester.email)
    ar_id = client.post(
        "/ars", json=_sample_ar_body(budget_no="7100-01", this_application="9327.10")
    ).json()["id"]
    _submit_for_approval(client, ar_id)

    _login_as(client, fa.email)
    blocked = client.post(f"/ars/{ar_id}/fa-acknowledge", json={"force": False})
    assert blocked.status_code == 409

    forced = client.post(f"/ars/{ar_id}/fa-acknowledge", json={"force": True})
    assert forced.status_code == 200
    body = forced.json()
    assert body["budget_overridden"] is True
    assert Decimal(str(body["budget_deducted_amount"])) == Decimal("9327.10")


# ───────── Revise ถูกจำกัดเฉพาะ Rejected เท่านั้น (Correction 2026-09-10) ─────────
# เดิม "Revise เมื่อไหร่ก็ได้หลัง Finalized" (แม้ Approved+หักงบไปแล้วจริง) — ผู้ใช้แจ้ง
# Business Rule ใหม่ 2026-09-10 ว่า Revise ได้เฉพาะฉบับที่ถูก Reject มาเท่านั้น จึงต้อง
# บล็อก AR ที่ Approved+FA Acknowledge แล้วไม่ให้ Revise อีกต่อไป (ไม่มี Endpoint ไหน
# Reject AR ที่ Approved ไปแล้วได้อีก — budget_workflow.refund_on_revise() จึงกลาย
# เป็น Path ที่ Route Layer ปัจจุบันเรียกไม่ถึงอีกแล้ว แต่ยังเก็บไว้เผื่อมี Endpoint
# "Reopen" ในอนาคต — ไม่ใช่ Bug ของ Revise Gate ปัจจุบัน)
def test_revise_blocked_after_approved_and_fa_acknowledged(
    client: TestClient, plain_user: User, db_session: Session
):
    fa = _make_user(db_session, name="FA Refund", email="farefund@example.com", is_fa=True)
    master = _make_budget_master(db_session, department="Production", account_code="5100-01")

    _login_as(client, plain_user.email, "plainpass123")
    ar_id = client.post("/ars", json=_sample_ar_body()).json()["id"]
    _submit_for_approval(client, ar_id)  # ไม่มี Level Setup ให้ Production ในเทสนี้ -> ตรงไป FA เลย

    _login_as(client, fa.email)
    ack = client.post(f"/ars/{ar_id}/fa-acknowledge", json={"force": False})
    assert ack.status_code == 200
    assert ack.json()["budget_approval_status"] == "approved"
    db_session.refresh(master)
    assert master.used_amount == Decimal("9327.10")

    # Approved แล้ว (ไม่ใช่ Rejected) -> Revise ต้องถูกบล็อก ยอดงบไม่ถูกคืนเพราะไม่ได้ Revise
    _login_as(client, plain_user.email, "plainpass123")
    revised = client.post(f"/ars/{ar_id}/revise")
    assert revised.status_code == 409

    db_session.refresh(master)
    assert master.used_amount == Decimal("9327.10")  # ยอดยังไม่ถูกคืน


# ───────── PATCH ระหว่างรออนุมัติ -> ยกเลิกคำขออัตโนมัติ (Correction 2026-09-10) ─────────
def test_patch_while_pending_cancels_approval_and_requires_resubmit(
    client: TestClient, plain_user: User, db_session: Session
):
    manager = _make_user(
        db_session, name="Manager Edit", email="mgredit@example.com", department="Production"
    )
    _make_level(
        db_session,
        department="Production",
        level_no=1,
        level_name="Manager",
        approver_user_id=manager.id,
    )

    _login_as(client, plain_user.email, "plainpass123")
    ar_id = client.post("/ars", json=_sample_ar_body()).json()["id"]
    submitted = _submit_for_approval(client, ar_id)
    assert submitted["status"] == "draft"  # มี Level ให้รอ -> ยังไม่ Finalize
    assert submitted["budget_approval_status"] == "pending"

    # แก้ไขระหว่างรออนุมัติ Level 1 อยู่ -> ต้องแก้ได้ (ยัง Draft) แต่ยกเลิกคำขอที่ค้างอยู่
    edited = client.patch(f"/ars/{ar_id}", json=_sample_ar_body(subject="แก้ไขยอดแล้ว"))
    assert edited.status_code == 200, edited.text
    body = edited.json()
    assert body["subject"] == "แก้ไขยอดแล้ว"
    assert body["budget_approval_status"] == "not_submitted"
    assert body["current_approval_level"] is None

    history = client.get(f"/ars/{ar_id}/history").json()
    actions = [log["action"] for log in history]
    assert "ar.approval_cancelled_by_edit" in actions

    # Level 1 เดิมกดอนุมัติไม่ได้อีกแล้วเพราะ current_approval_level ถูกล้างไปแล้ว
    _login_as(client, manager.email)
    stale = client.post(f"/ars/{ar_id}/approve-level")
    assert stale.status_code == 409

    # ต้องกด "ส่งขออนุมัติ" ใหม่ถึงจะเข้าคิว Level 1 อีกครั้ง
    _login_as(client, plain_user.email, "plainpass123")
    resubmitted = _submit_for_approval(client, ar_id)
    assert resubmitted["budget_approval_status"] == "pending"
    assert resubmitted["current_approval_level"] == 1


def test_submit_for_approval_already_pending_rejected(
    client: TestClient, plain_user: User, db_session: Session
):
    manager = _make_user(
        db_session, name="Manager Dup", email="mgrdup@example.com", department="Production"
    )
    _make_level(
        db_session,
        department="Production",
        level_no=1,
        level_name="Manager",
        approver_user_id=manager.id,
    )

    _login_as(client, plain_user.email, "plainpass123")
    ar_id = client.post("/ars", json=_sample_ar_body()).json()["id"]
    first = client.post(f"/ars/{ar_id}/submit-for-approval")
    assert first.status_code == 200
    assert first.json()["status"] == "draft"  # มี Level ให้รอ -> ยังไม่ Finalize

    # ยังเป็น Draft อยู่ (ไม่โดน Gate สถานะ) แต่ถูก Gate สถานะ Not Submitted แทน เพราะส่ง
    # ขออนุมัติไปแล้วครั้งนึง (Pending) — กด "ส่งขออนุมัติ" ซ้ำไม่ได้จนกว่าจะจบ Level หรือ
    # ถูก Reject/แก้ไข (ยกเลิกคำขอ) ก่อน
    second = client.post(f"/ars/{ar_id}/submit-for-approval")
    assert second.status_code == 409


# ───────────────────────── Excel Upload ─────────────────────────
def _build_xlsx_bytes(rows: list[list]) -> bytes:
    import io

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(
        [
            "budget_no",
            "department",
            "budget_type",
            "account_code",
            "period_start",
            "period_end",
            "budget_name",
            "budgeted_amount",
        ]
    )
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_excel_upload_requires_fa_or_admin(client: TestClient, plain_user: User):
    _login_as(client, plain_user.email, "plainpass123")
    content = _build_xlsx_bytes([])
    res = client.post(
        "/budget/upload",
        files={
            "file": (
                "budget.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert res.status_code == 403


def test_excel_upload_partial_success_and_reupload_merge(client: TestClient, db_session: Session):
    fa = _make_user(db_session, name="FA Upload", email="faupload@example.com", is_fa=True)
    _login_as(client, fa.email)

    content = _build_xlsx_bytes(
        [
            [
                "BGW001",
                "Warehouse",
                "expenses",
                "8100-01",
                "01/01/2026",
                "31/12/2026",
                "Warehouse Budget",
                50000,
            ],
            [
                "BGW002",
                "Warehouse",
                "expenses",
                "8100-02",
                "01/01/2026",
                "31/12/2026",
                None,
                -100,
            ],  # ผิด: ติดลบ
            [
                "BGW003",
                "Warehouse",
                "assets",
                "8200-01",
                "31/12/2026",
                "01/01/2026",
                None,
                1000,
            ],  # ผิด: start > end
        ]
    )
    res = client.post(
        "/budget/upload",
        files={
            "file": (
                "budget.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert res.status_code == 201, res.text
    body = res.json()
    assert body["total_rows"] == 3
    assert body["success_rows"] == 1
    assert body["error_rows"] == 2

    listed = client.get("/budget", params={"department": "Warehouse"}).json()
    assert len(listed) == 1
    assert listed[0]["budgeted_amount"] == "50000.00" or Decimal(
        str(listed[0]["budgeted_amount"])
    ) == Decimal("50000.00")

    # Re-upload budget_no เดิม แต่เปลี่ยน budgeted_amount — ต้อง Merge/Update แถวเดิม
    # ไม่สร้างซ้ำ (Key จับคู่คือ budget_no เดี่ยว ไม่ใช่ account_code+period อีกต่อไป)
    content2 = _build_xlsx_bytes(
        [
            [
                "BGW001",
                "Warehouse",
                "expenses",
                "8100-01",
                "01/01/2026",
                "31/12/2026",
                "Warehouse Budget 2",
                60000,
            ]
        ]
    )
    res2 = client.post(
        "/budget/upload",
        files={
            "file": (
                "budget2.xlsx",
                content2,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert res2.status_code == 201
    assert res2.json()["success_rows"] == 1

    listed2 = client.get(
        "/budget", params={"department": "Warehouse", "account_code": "8100-01"}
    ).json()
    assert len(listed2) == 1
    assert Decimal(str(listed2[0]["budgeted_amount"])) == Decimal("60000.00")


# ───────────────────────── Regression: account_code ซ้ำกันได้หลาย budget_no ─────────────────────────
# ตัวอย่างจริงที่ผู้ใช้ส่งมา 2026-09-09 (สาเหตุที่แก้ Design v4.1): account_code เดียวกัน
# (เช่น 5100) มีได้หลายรายการงบ คนละ budget_no (BG0001, BG0002) — ก่อนแก้ Unique
# Constraint เดิมผูกกับ account_code ทำให้แถวที่ 2 Insert ไม่ผ่าน
def test_excel_upload_same_account_code_different_budget_no_both_succeed(
    client: TestClient, db_session: Session
):
    fa = _make_user(db_session, name="FA Dup Code", email="fadupcode@example.com", is_fa=True)
    _login_as(client, fa.email)

    content = _build_xlsx_bytes(
        [
            [
                "BG0001",
                "Production",
                "expenses",
                "5100",
                "01/01/2026",
                "31/12/2026",
                "ค่าใช้จ่ายซ่อมบำรุงเครื่องจักร 300T",
                500000,
            ],
            [
                "BG0002",
                "Production",
                "expenses",
                "5100",
                "01/01/2026",
                "31/12/2026",
                "ค่าใช้จ่ายซ่อมบำรุงเครื่องจักร 1000T",
                100000,
            ],
        ]
    )
    res = client.post(
        "/budget/upload",
        files={
            "file": (
                "budget.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert res.status_code == 201, res.text
    body = res.json()
    assert body["total_rows"] == 2
    assert body["success_rows"] == 2, body
    assert body["error_rows"] == 0

    listed = client.get(
        "/budget", params={"department": "Production", "account_code": "5100"}
    ).json()
    assert len(listed) == 2
    assert {row["budget_no"] for row in listed} == {"BG0001", "BG0002"}


def test_excel_upload_duplicate_budget_no_in_same_file_rejected(
    client: TestClient, db_session: Session
):
    fa = _make_user(db_session, name="FA Dup No", email="fadupno@example.com", is_fa=True)
    _login_as(client, fa.email)

    content = _build_xlsx_bytes(
        [
            ["BG9001", "Production", "expenses", "5100", "01/01/2026", "31/12/2026", None, 1000],
            ["BG9001", "Production", "expenses", "1600", "01/01/2026", "31/12/2026", None, 2000],
        ]
    )
    res = client.post(
        "/budget/upload",
        files={
            "file": (
                "budget.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert res.status_code == 201, res.text
    body = res.json()
    assert body["success_rows"] == 1
    assert body["error_rows"] == 1
    assert "ซ้ำ" in [r["message"] for r in body["rows"] if not r["ok"]][0]


# ───────────────────────── Comment ตอนอนุมัติ (Phase A, 2026-09-10) ─────────────────────────
def test_approve_level_comment_is_optional_and_shown_in_progress_and_pdf(
    client: TestClient, plain_user: User, db_session: Session
):
    """Correction 2026-09-10 (Phase A): approve-level รับ comment ไม่บังคับแล้ว — เก็บลง
    คอลัมน์ reason เดิม (Reuse ร่วมกับเหตุผลปฏิเสธ) โชว์ทั้งใน approval-progress และตาราง
    Authority ของ PDF"""
    manager = _make_user(
        db_session, name="Manager Cmt", email="mgrcmt@example.com", department="Production"
    )
    _make_level(
        db_session,
        department="Production",
        level_no=1,
        level_name="Manager",
        approver_user_id=manager.id,
    )
    _make_budget_master(db_session)

    _login_as(client, plain_user.email, "plainpass123")
    ar_id = client.post("/ars", json=_sample_ar_body()).json()["id"]
    _submit_for_approval(client, ar_id)

    _login_as(client, manager.email)
    # ไม่ใส่ Comment เลย (Body ว่างเปล่าได้) — ยังทำงานได้ปกติเหมือนเดิมก่อน Phase A
    no_comment = client.post(f"/ars/{ar_id}/approve-level", json={})
    assert no_comment.status_code == 200, no_comment.text

    progress = client.get(f"/ars/{ar_id}/approval-progress").json()
    assert progress[0]["reason"] is None


def test_approve_level_with_comment_stored(
    client: TestClient, plain_user: User, db_session: Session
):
    manager = _make_user(
        db_session, name="Manager Note", email="mgrnote@example.com", department="Production"
    )
    _make_level(
        db_session,
        department="Production",
        level_no=1,
        level_name="Manager",
        approver_user_id=manager.id,
    )
    fa = _make_user(db_session, name="FA Note", email="fanote@example.com", is_fa=True)
    _make_budget_master(db_session)

    _login_as(client, plain_user.email, "plainpass123")
    ar_id = client.post("/ars", json=_sample_ar_body()).json()["id"]
    _submit_for_approval(client, ar_id)

    _login_as(client, manager.email)
    approved = client.post(f"/ars/{ar_id}/approve-level", json={"comment": "เร่งด่วน อนุมัติให้เลย"})
    assert approved.status_code == 200, approved.text

    progress = client.get(f"/ars/{ar_id}/approval-progress").json()
    assert progress[0]["reason"] == "เร่งด่วน อนุมัติให้เลย"

    _login_as(client, fa.email)
    ack = client.post(
        f"/ars/{ar_id}/fa-acknowledge", json={"force": False, "comment": "หักงบเรียบร้อย"}
    )
    assert ack.status_code == 200, ack.text

    progress2 = client.get(f"/ars/{ar_id}/approval-progress").json()
    assert progress2[-1]["reason"] == "หักงบเรียบร้อย"

    # PDF Auto-fill: Comment ตอนอนุมัติต้องโผล่ในช่อง Comments ของตาราง Authority
    ar_obj = db_session.get(ApprovalRequest, ar_id)
    html = render_ar_html(db_session, ar_obj)
    assert "เร่งด่วน อนุมัติให้เลย" in html


# ───────────────── Add/Edit Budget ด้วยตัวเอง (Correction 2026-09-11) ─────────────────
def _sample_create_body(**overrides) -> dict:
    body = {
        "budget_no": "BGM001",
        "department": "Manual Dept",
        "budget_type": "expenses",
        "account_code": "5900",
        "period_start": "2026-01-01",
        "period_end": "2026-12-31",
        "budgeted_amount": "20000.00",
        "budget_name": "เพิ่มเอง",
    }
    body.update(overrides)
    return body


def test_create_budget_master_requires_fa_or_admin(client: TestClient, plain_user: User):
    _login_as(client, plain_user.email, "plainpass123")
    res = client.post("/budget", json=_sample_create_body())
    assert res.status_code == 403


def test_create_budget_master_success(client: TestClient, db_session: Session):
    fa = _make_user(db_session, name="FA Add", email="faadd@example.com", is_fa=True)
    _login_as(client, fa.email)

    res = client.post("/budget", json=_sample_create_body())
    assert res.status_code == 201, res.text
    body = res.json()
    assert body["budget_no"] == "BGM001"
    assert Decimal(str(body["budgeted_amount"])) == Decimal("20000.00")
    assert Decimal(str(body["used_amount"])) == Decimal("0")
    assert Decimal(str(body["balance"])) == Decimal("20000.00")

    listed = client.get("/budget", params={"budget_no": "BGM001"}).json()
    assert len(listed) == 1
    assert listed[0]["department"] == "Manual Dept"


def test_create_budget_master_duplicate_budget_no_rejected(client: TestClient, db_session: Session):
    _make_budget_master(db_session, budget_no="BGM002")
    fa = _make_user(db_session, name="FA Dup Add", email="fadupadd@example.com", is_fa=True)
    _login_as(client, fa.email)

    res = client.post("/budget", json=_sample_create_body(budget_no="BGM002"))
    assert res.status_code == 409


def test_create_budget_master_invalid_period_rejected(client: TestClient, db_session: Session):
    fa = _make_user(db_session, name="FA Bad Period", email="fabadperiod@example.com", is_fa=True)
    _login_as(client, fa.email)

    res = client.post(
        "/budget",
        json=_sample_create_body(period_start="2026-12-31", period_end="2026-01-01"),
    )
    assert res.status_code == 400


def test_create_budget_master_negative_amount_rejected(client: TestClient, db_session: Session):
    fa = _make_user(db_session, name="FA Negative", email="fanegative@example.com", is_fa=True)
    _login_as(client, fa.email)

    res = client.post("/budget", json=_sample_create_body(budgeted_amount="-1"))
    assert res.status_code == 422


def test_update_budget_master_requires_fa_or_admin(
    client: TestClient, plain_user: User, db_session: Session
):
    row = _make_budget_master(db_session, budget_no="BGM003")
    _login_as(client, plain_user.email, "plainpass123")
    res = client.patch(f"/budget/{row.id}", json={"department": "Hacked"})
    assert res.status_code == 403


def test_update_budget_master_success(client: TestClient, db_session: Session):
    row = _make_budget_master(db_session, budget_no="BGM004", department="Old Dept")
    fa = _make_user(db_session, name="FA Edit", email="faedit@example.com", is_fa=True)
    _login_as(client, fa.email)

    res = client.patch(
        f"/budget/{row.id}",
        json={"department": "New Dept", "budgeted_amount": "99999.00"},
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["department"] == "New Dept"
    assert Decimal(str(body["budgeted_amount"])) == Decimal("99999.00")
    # budget_no ไม่ถูกส่งมาแก้ ต้องเหมือนเดิม
    assert body["budget_no"] == "BGM004"


def test_update_budget_master_not_found(client: TestClient, db_session: Session):
    fa = _make_user(db_session, name="FA 404", email="fa404@example.com", is_fa=True)
    _login_as(client, fa.email)
    res = client.patch("/budget/999999", json={"department": "X"})
    assert res.status_code == 404


def test_update_budget_master_invalid_period_rejected(client: TestClient, db_session: Session):
    row = _make_budget_master(db_session, budget_no="BGM005")
    fa = _make_user(db_session, name="FA Bad Period2", email="fabadperiod2@example.com", is_fa=True)
    _login_as(client, fa.email)

    res = client.patch(
        f"/budget/{row.id}",
        json={"period_start": "2026-12-31", "period_end": "2026-01-01"},
    )
    assert res.status_code == 400


def test_update_budget_master_cannot_change_budget_no_or_used_amount(
    client: TestClient, db_session: Session
):
    row = _make_budget_master(db_session, budget_no="BGM006", used_amount=Decimal("500.00"))
    fa = _make_user(db_session, name="FA Protect", email="faprotect@example.com", is_fa=True)
    _login_as(client, fa.email)

    # Schema BudgetMasterUpdate ไม่มี Field budget_no/used_amount เลย — ส่งมาก็ถูก
    # Pydantic ตัดทิ้งเงียบๆ ไม่มีผลอะไร
    res = client.patch(
        f"/budget/{row.id}",
        json={
            "budget_no": "SHOULD-NOT-CHANGE",
            "used_amount": "999999.00",
            "department": "Still Works",
        },
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["budget_no"] == "BGM006"
    assert Decimal(str(body["used_amount"])) == Decimal("500.00")
    assert body["department"] == "Still Works"


# ───────────────────────── Export Budget ─────────────────────────
def test_export_budget_master_requires_fa_or_admin(client: TestClient, plain_user: User):
    _login_as(client, plain_user.email, "plainpass123")
    res = client.get("/budget/export")
    assert res.status_code == 403


def test_export_budget_master_success(client: TestClient, db_session: Session):
    _make_budget_master(
        db_session,
        budget_no="BGX001",
        department="Export Dept",
        budgeted_amount=Decimal("10000.00"),
        used_amount=Decimal("1500.00"),
    )
    fa = _make_user(db_session, name="FA Export", email="faexport@example.com", is_fa=True)
    _login_as(client, fa.email)

    res = client.get("/budget/export")
    assert res.status_code == 200
    assert (
        res.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in res.headers["content-disposition"]

    import io

    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    assert header[:8] == (
        "budget_no",
        "department",
        "budget_type",
        "account_code",
        "period_start",
        "period_end",
        "budgeted_amount",
        "budget_name",
    )
    assert header[8] == "used_amount"
    assert header[9] == "balance"

    data_rows = [r for r in rows[1:] if r[0] == "BGX001"]
    assert len(data_rows) == 1
    row = data_rows[0]
    assert row[1] == "Export Dept"
    assert row[6] == 10000.0
    assert row[8] == 1500.0
    assert row[9] == 8500.0
