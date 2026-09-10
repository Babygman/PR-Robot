"""Test เอกสารแนบของ Approval Request (AR Attachments, Phase A — Correction 2026-09-10)

Pattern การ Login/สร้าง User/Level เดียวกับ tests/test_budget_control.py — ครอบคลุม
Permission ตาม Docstring ของ budget_workflow.can_upload_attachment: เจ้าของ AR/Admin
แนบ+ลบได้เสมอ (ลบต้องเป็นคนอัปโหลดเองด้วยถ้าไม่ใช่เจ้าของ/Admin), ผู้อนุมัติ Level
ปัจจุบัน/FA แนบได้เฉพาะช่วงที่กำลังรอตัดสินใจอยู่จริงเท่านั้น
"""

from __future__ import annotations

import io

from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy.orm import Session

from app.models import User
from tests.test_approval_requests import _sample_ar_body
from tests.test_budget_control import _login_as, _make_level, _make_user, _submit_for_approval

_FAKE_PDF = b"%PDF-1.4 fake content for testing\n"


def _make_real_xlsx_bytes() -> bytes:
    """Comment 4 (2026-09-10 — xlsx-preview): ต้องเป็นไฟล์ .xlsx จริง (ไม่ใช่ Bytes ปลอม
    เหมือน test_upload_accepts_excel_and_image) เพราะ Endpoint นี้เปิดอ่านด้วย openpyxl
    จริงๆ ไม่ได้แค่เช็ค Content-Type ตอน Upload เท่านั้น"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Budget"
    sheet.append(["Item", "Qty", "Amount"])
    sheet.append(["pc", 1, 44040])
    sheet.append(["mouse", 2, 250])
    buf = io.BytesIO()
    workbook.save(buf)
    return buf.getvalue()


def _make_multi_sheet_xlsx_bytes() -> bytes:
    """แก้ไขเพิ่ม (2026-09-10): ผู้ใช้แจ้งว่าไฟล์ Excel จริงมีหลาย Tab (Sheet) ต้อง Preview
    เห็นครบทุก Tab ไม่ใช่แค่ Tab แรก — จำลองไฟล์ 3 Sheet เหมือนไฟล์จริงที่ผู้ใช้ส่งมา"""
    workbook = Workbook()
    sheet1 = workbook.active
    sheet1.title = "Budget Master"
    sheet1.append(["Department", "Budget No"])
    sheet1.append(["Production", "BG0001"])
    sheet2 = workbook.create_sheet("Approval")
    sheet2.append(["Level", "Approver"])
    sheet2.append(["Manager", "Somchai"])
    sheet3 = workbook.create_sheet("Notes")
    sheet3.append(["Remark"])
    buf = io.BytesIO()
    workbook.save(buf)
    return buf.getvalue()


def _login(client: TestClient) -> None:
    client.post("/auth/login", json={"email": "plain@example.com", "password": "plainpass123"})


def test_upload_requires_login(client: TestClient):
    res = client.post(
        "/ars/1/attachments",
        files={"file": ("q.pdf", io.BytesIO(_FAKE_PDF), "application/pdf")},
    )
    assert res.status_code == 401


def test_owner_can_upload_list_download_pdf_attachment(client: TestClient, plain_user: User):
    _login(client)
    ar_id = client.post("/ars", json=_sample_ar_body()).json()["id"]

    up = client.post(
        f"/ars/{ar_id}/attachments",
        files={"file": ("quotation.pdf", io.BytesIO(_FAKE_PDF), "application/pdf")},
    )
    assert up.status_code == 201, up.text
    body = up.json()
    assert body["file_name"] == "quotation.pdf"
    assert body["uploaded_by_id"] == plain_user.id
    assert body["file_size"] == len(_FAKE_PDF)

    listed = client.get(f"/ars/{ar_id}/attachments").json()
    assert len(listed) == 1
    assert listed[0]["id"] == body["id"]

    dl = client.get(f"/ars/{ar_id}/attachments/{body['id']}/download")
    assert dl.status_code == 200
    assert dl.content == _FAKE_PDF


def test_upload_accepts_excel_and_image(client: TestClient, plain_user: User):
    _login(client)
    ar_id = client.post("/ars", json=_sample_ar_body()).json()["id"]

    xlsx = client.post(
        f"/ars/{ar_id}/attachments",
        files={
            "file": (
                "budget.xlsx",
                io.BytesIO(b"fake xlsx bytes"),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert xlsx.status_code == 201, xlsx.text

    png = client.post(
        f"/ars/{ar_id}/attachments",
        files={"file": ("photo.png", io.BytesIO(b"fake png bytes"), "image/png")},
    )
    assert png.status_code == 201, png.text

    assert len(client.get(f"/ars/{ar_id}/attachments").json()) == 2


def test_upload_rejects_unsupported_content_type(client: TestClient, plain_user: User):
    _login(client)
    ar_id = client.post("/ars", json=_sample_ar_body()).json()["id"]
    res = client.post(
        f"/ars/{ar_id}/attachments",
        files={"file": ("archive.zip", io.BytesIO(b"PK\x03\x04fake"), "application/zip")},
    )
    assert res.status_code == 415


def test_delete_by_uploader_ok_by_unrelated_user_forbidden_by_admin_ok(
    client: TestClient, plain_user: User, admin_user: User, db_session: Session
):
    other = _make_user(
        db_session, name="Unrelated", email="unrelated@example.com", department="Production"
    )

    _login(client)
    ar_id = client.post("/ars", json=_sample_ar_body()).json()["id"]
    att_id = client.post(
        f"/ars/{ar_id}/attachments",
        files={"file": ("q.pdf", io.BytesIO(_FAKE_PDF), "application/pdf")},
    ).json()["id"]

    _login_as(client, other.email)
    forbidden = client.delete(f"/ars/{ar_id}/attachments/{att_id}")
    assert forbidden.status_code == 403

    _login_as(client, admin_user.email, "adminpass123")
    ok = client.delete(f"/ars/{ar_id}/attachments/{att_id}")
    assert ok.status_code == 204

    _login(client)
    assert client.get(f"/ars/{ar_id}/attachments").json() == []


def test_xlsx_preview_returns_rows_for_real_xlsx(client: TestClient, plain_user: User):
    _login(client)
    ar_id = client.post("/ars", json=_sample_ar_body()).json()["id"]

    att_id = client.post(
        f"/ars/{ar_id}/attachments",
        files={
            "file": (
                "budget.xlsx",
                io.BytesIO(_make_real_xlsx_bytes()),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    ).json()["id"]

    res = client.get(f"/ars/{ar_id}/attachments/{att_id}/xlsx-preview")
    assert res.status_code == 200, res.text
    body = res.json()
    assert len(body["sheets"]) == 1
    sheet = body["sheets"][0]
    assert sheet["sheet_name"] == "Budget"
    assert sheet["truncated"] is False
    assert sheet["rows"] == [
        ["Item", "Qty", "Amount"],
        ["pc", 1, 44040],
        ["mouse", 2, 250],
    ]


def test_xlsx_preview_returns_all_sheets(client: TestClient, plain_user: User):
    _login(client)
    ar_id = client.post("/ars", json=_sample_ar_body()).json()["id"]

    att_id = client.post(
        f"/ars/{ar_id}/attachments",
        files={
            "file": (
                "budget_master_template and approval.xlsx",
                io.BytesIO(_make_multi_sheet_xlsx_bytes()),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    ).json()["id"]

    res = client.get(f"/ars/{ar_id}/attachments/{att_id}/xlsx-preview")
    assert res.status_code == 200, res.text
    sheets = res.json()["sheets"]
    assert [s["sheet_name"] for s in sheets] == ["Budget Master", "Approval", "Notes"]
    assert sheets[0]["rows"] == [["Department", "Budget No"], ["Production", "BG0001"]]
    assert sheets[1]["rows"] == [["Level", "Approver"], ["Manager", "Somchai"]]
    assert sheets[2]["rows"] == [["Remark"]]


def test_xlsx_preview_rejects_unreadable_file(client: TestClient, plain_user: User):
    _login(client)
    ar_id = client.post("/ars", json=_sample_ar_body()).json()["id"]

    att_id = client.post(
        f"/ars/{ar_id}/attachments",
        files={
            "file": (
                "budget.xlsx",
                io.BytesIO(b"not actually a real xlsx file"),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    ).json()["id"]

    res = client.get(f"/ars/{ar_id}/attachments/{att_id}/xlsx-preview")
    assert res.status_code == 422


def test_current_level_approver_can_upload_while_pending_other_approver_cannot(
    client: TestClient, plain_user: User, db_session: Session
):
    manager = _make_user(
        db_session, name="Manager Att", email="mgratt@example.com", department="Production"
    )
    gm = _make_user(db_session, name="GM Att", email="gmatt@example.com", department="Production")
    _make_level(
        db_session,
        department="Production",
        level_no=1,
        level_name="Manager",
        approver_user_id=manager.id,
    )
    _make_level(
        db_session,
        department="Production",
        level_no=2,
        level_name="General Manager",
        approver_user_id=gm.id,
    )

    _login(client)
    ar_id = client.post("/ars", json=_sample_ar_body()).json()["id"]
    _submit_for_approval(
        client, ar_id
    )  # budget_approval_status == pending, current_approval_level == 1

    # Level 1 (manager) กำลังรออยู่ — แนบเอกสารเพิ่มได้ระหว่างรอตัดสินใจ
    _login_as(client, manager.email)
    ok = client.post(
        f"/ars/{ar_id}/attachments",
        files={"file": ("extra.pdf", io.BytesIO(_FAKE_PDF), "application/pdf")},
    )
    assert ok.status_code == 201, ok.text

    # Level 2 (gm) ยังไม่ถึงคิว — แนบไม่ได้
    _login_as(client, gm.email)
    forbidden = client.post(
        f"/ars/{ar_id}/attachments",
        files={"file": ("extra2.pdf", io.BytesIO(_FAKE_PDF), "application/pdf")},
    )
    assert forbidden.status_code == 403


def test_fa_can_upload_only_during_pending_fa_acknowledge(
    client: TestClient, plain_user: User, db_session: Session
):
    fa = _make_user(db_session, name="FA Att", email="faatt@example.com", is_fa=True)

    _login(client)
    ar_id = client.post("/ars", json=_sample_ar_body()).json()["id"]

    # ยังไม่ Submit (not_submitted) — FA ยังแนบไม่ได้ (ไม่ใช่เจ้าของ/Admin)
    _login_as(client, fa.email)
    too_early = client.post(
        f"/ars/{ar_id}/attachments",
        files={"file": ("early.pdf", io.BytesIO(_FAKE_PDF), "application/pdf")},
    )
    assert too_early.status_code == 403

    _login(client)
    _submit_for_approval(client, ar_id)  # แผนก Production ไม่มี Level -> ข้ามตรงไป FA ทันที

    _login_as(client, fa.email)
    ok = client.post(
        f"/ars/{ar_id}/attachments",
        files={"file": ("beforeack.pdf", io.BytesIO(_FAKE_PDF), "application/pdf")},
    )
    assert ok.status_code == 201, ok.text
