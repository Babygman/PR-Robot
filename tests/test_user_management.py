"""Test User Management Redesign (2026-09-11): division field, Delete User (ลบได้
เฉพาะไม่มีประวัติ), Excel Download/Upload — Pattern เดียวกับ tests/test_budget_control.py
สำหรับส่วน Excel"""

from __future__ import annotations

import io
from datetime import date

import openpyxl
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.core.security import hash_password
from app.models import ApprovalRequest, User

_PASSWORD = "password123456"


def _make_admin(db_session: Session, **kwargs) -> User:
    kwargs.setdefault("password_hash", hash_password(_PASSWORD))
    kwargs.setdefault("name", "Admin User")
    kwargs.setdefault("email", "admin2@example.com")
    kwargs.setdefault("is_admin", True)
    user = User(**kwargs)
    db_session.add(user)
    db_session.commit()
    db_session.refresh(user)
    return user


def _login_as(client: TestClient, email: str, password: str = _PASSWORD) -> None:
    res = client.post("/auth/login", json={"email": email, "password": password})
    assert res.status_code == 200, res.text


def test_create_user_stores_division(client: TestClient, db_session: Session):
    _make_admin(db_session)
    _login_as(client, "admin2@example.com")
    res = client.post(
        "/users",
        json={
            "name": "Somchai",
            "email": "somchai@example.com",
            "password": "somchaipass123",
            "department": "Production",
            "division": "Manufacturing",
            "position": "Officer",
        },
    )
    assert res.status_code == 201, res.text
    body = res.json()
    assert body["division"] == "Manufacturing"


def test_update_user_division(client: TestClient, db_session: Session):
    admin = _make_admin(db_session)
    target = User(
        name="Target",
        email="target@example.com",
        password_hash=hash_password(_PASSWORD),
    )
    db_session.add(target)
    db_session.commit()
    db_session.refresh(target)

    _login_as(client, admin.email)
    res = client.patch(f"/users/{target.id}", json={"division": "Quality Assurance"})
    assert res.status_code == 200, res.text
    assert res.json()["division"] == "Quality Assurance"


def test_delete_user_without_history_succeeds(client: TestClient, db_session: Session):
    admin = _make_admin(db_session)
    fresh = User(
        name="Fresh User",
        email="fresh@example.com",
        password_hash=hash_password(_PASSWORD),
    )
    db_session.add(fresh)
    db_session.commit()
    db_session.refresh(fresh)

    _login_as(client, admin.email)
    res = client.delete(f"/users/{fresh.id}")
    assert res.status_code == 204, res.text
    assert db_session.get(User, fresh.id) is None


def test_delete_user_with_ar_history_is_blocked(client: TestClient, db_session: Session):
    admin = _make_admin(db_session)
    requester = User(
        name="Requester",
        email="requester@example.com",
        password_hash=hash_password(_PASSWORD),
        department="Production",
        can_view_ar=True,
    )
    db_session.add(requester)
    db_session.commit()
    db_session.refresh(requester)

    ar = ApprovalRequest(
        ar_no=1,
        status="draft",
        requested_by_id=requester.id,
        application_date=date(2026, 9, 1),
        subject="Test",
        budget_type="expenses",
    )
    db_session.add(ar)
    db_session.commit()

    _login_as(client, admin.email)
    res = client.delete(f"/users/{requester.id}")
    assert res.status_code == 409, res.text
    detail = res.json()["detail"]
    assert "ประวัติ" in detail or "PR" in detail or "Approval" in detail
    # ยัง Active/Inactive Toggle ได้ปกติแทน
    assert db_session.get(User, requester.id) is not None


def test_delete_user_not_found(client: TestClient, db_session: Session):
    admin = _make_admin(db_session)
    _login_as(client, admin.email)
    res = client.delete("/users/999999")
    assert res.status_code == 404


def test_delete_user_requires_admin(client: TestClient, plain_user: User):
    _login_as(client, plain_user.email, "plainpass123")
    res = client.delete(f"/users/{plain_user.id}")
    assert res.status_code == 403


def _build_users_xlsx_bytes(rows: list[list]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(
        [
            "name",
            "email",
            "password",
            "department",
            "division",
            "position",
            "is_admin",
            "is_fa",
            "can_view_approvals",
            "can_view_pr",
            "can_view_ar",
            "can_view_all_pr",
            "can_view_all_ar",
            "is_active",
        ]
    )
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_export_users_requires_admin(client: TestClient, plain_user: User):
    _login_as(client, plain_user.email, "plainpass123")
    res = client.get("/users/export")
    assert res.status_code == 403


def test_export_users_success(client: TestClient, db_session: Session):
    admin = _make_admin(db_session)
    _login_as(client, admin.email)
    res = client.get("/users/export")
    assert res.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb.active
    header = [c.value for c in next(ws.iter_rows(max_row=1))]
    assert "email" in header
    assert "division" in header
    assert "password" not in header  # ไม่ Export รหัสผ่านออกมาด้วยเหตุผลความปลอดภัย


def test_upload_users_creates_and_updates(client: TestClient, db_session: Session):
    admin = _make_admin(db_session)
    existing = User(
        name="Old Name",
        email="upsert@example.com",
        password_hash=hash_password(_PASSWORD),
    )
    db_session.add(existing)
    db_session.commit()
    db_session.refresh(existing)

    _login_as(client, admin.email)
    content = _build_users_xlsx_bytes(
        [
            [
                "New Guy",
                "newguy@example.com",
                "newguypass123",
                "Warehouse",
                "Logistics",
                "Officer",
                False,
                False,
                False,
                True,
                False,
                False,
                False,
                True,
            ],
            [
                "Updated Name",
                "upsert@example.com",
                "",  # แถว Update — ไม่ใช้คอลัมน์ password
                "Sales",
                "",
                "",
                False,
                False,
                False,
                True,
                True,
                False,
                False,
                True,
            ],
            [
                "Missing Password",
                "nopass@example.com",
                "",
                "",
                "",
                "",
                False,
                False,
                False,
                False,
                False,
                False,
                False,
                True,
            ],  # ผิด: User ใหม่ต้องมี password
        ]
    )
    res = client.post(
        "/users/upload",
        files={
            "file": (
                "users.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert res.status_code == 201, res.text
    body = res.json()
    assert body["total_rows"] == 3
    assert body["success_rows"] == 2
    assert body["error_rows"] == 1

    new_user = db_session.query(User).filter(User.email == "newguy@example.com").first()
    assert new_user is not None
    assert new_user.division == "Logistics"

    db_session.refresh(existing)
    assert existing.name == "Updated Name"
    assert existing.department == "Sales"
    # รหัสผ่านเดิมต้องไม่ถูกแตะต้อง (Upload ไม่ใช้คอลัมน์ password สำหรับแถว Update)
    from app.core.security import verify_password

    assert verify_password(_PASSWORD, existing.password_hash)


def test_upload_users_requires_admin(client: TestClient, plain_user: User):
    _login_as(client, plain_user.email, "plainpass123")
    content = _build_users_xlsx_bytes([])
    res = client.post(
        "/users/upload",
        files={
            "file": (
                "users.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert res.status_code == 403
