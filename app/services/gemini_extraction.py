"""เรียก Google Gemini API สกัดข้อมูลจากเอกสารต้นทาง (ใบเสนอราคา/ใบรับของ)

ใช้ SDK ใหม่ google-genai (google-generativeai รุ่นเก่าเลิกใช้แล้วตั้งแต่ปี 2025)
Model กำหนดผ่าน .env (GEMINI_MODEL) ไม่ Hardcode เพื่อให้ปรับได้ตาม Quota จริงของ API Key

Client ถูกออกแบบให้ Inject แทนที่ได้ (Dependency Injection) เพื่อให้ Unit Test
ไม่ต้องยิง Network จริงไปหา Gemini — ทดสอบด้วย Fake Client แทน
"""
from __future__ import annotations

import csv
import time
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path

import docx
import httpx
import openpyxl
from google import genai
from google.genai import errors as genai_errors
from google.genai import types

from app.core.config import settings
from app.schemas.source_document import ExtractionResult

_EXTRACTION_PROMPT = """\
คุณคือระบบสกัดข้อมูลจากเอกสารจัดซื้อของบริษัท Sunstar Chemical (Thailand) ให้แม่นยำ
ที่สุดเท่าที่ทำได้ เอกสารอาจเป็นใบเสนอราคา ใบยืมสินค้า ใบส่งสินค้า หรือใบรับของ — และมา
จาก Supplier ได้หลายรายที่มีรูปแบบตารางไม่เหมือนกันเลย (ตำแหน่งคอลัมน์ ชื่อหัวตาราง ภาษา
ที่ใช้ ต่างกันได้) ให้อ่านตามความหมายของแต่ละส่วนในเอกสารจริง ไม่ใช่ตำแหน่งตายตัว

อ่านเอกสารที่แนบมาแล้วดึงข้อมูลต่อไปนี้:
- detected_doc_type: ประเภทเอกสารที่คุณอ่านออกจากเนื้อหา ต้องเป็นค่าใดค่าหนึ่งนี้เท่านั้น
  "quotation" (ใบเสนอราคา), "receiving_note" (ใบรับของ), "borrow_note" (ใบยืมสินค้า),
  "delivery_note" (ใบส่งสินค้า), หรือ "other" ถ้าไม่แน่ใจ
- vendor_name: ชื่อผู้ขาย/ร้านค้า/Supplier
- document_no: เลขที่เอกสาร
- document_date: วันที่เอกสาร (รูปแบบ YYYY-MM-DD)
- items: รายการสินค้าทุกรายการ แยกฟิลด์ต่อไปนี้ให้ชัดเจนที่สุด (เอกสารส่วนใหญ่มีคอลัมน์
  แยกกันอยู่แล้ว):
  - product_name: ชื่อสินค้า (เช่น คอลัมน์ "PRODUCT NAME")
  - description: รายละเอียด/สเปกสินค้า (เช่น คอลัมน์ "PRODUCT DESCRIPTION" ขนาด/มิติ)
  - color: สี (ถ้ามีคอลัมน์แยก เช่น "COLOR")
  - quantity: จำนวน (ตัวเลขล้วนๆ ไม่รวมหน่วย)
  - unit: หน่วยนับ (เช่น Roll, ชิ้น, กก. — เอกสารบางแบบย่อไว้เช่น "R." ให้ขยายเป็นคำเต็ม
    ถ้ามั่นใจความหมาย เช่น "R." มักหมายถึง "Roll")
  - unit_price, amount: ราคาต่อหน่วย/ราคารวม ถ้ามี
- notes: หมายเหตุอื่นที่เกี่ยวข้องกับการจัดซื้อ
- ai_confidence: ประเมินความมั่นใจของคุณเองว่าอ่านข้อมูลถูกต้องครบถ้วนแค่ไหน (ตัวเลข 0.0-1.0)

กฎสำคัญ: ถ้าอ่านข้อมูลใดไม่ได้หรือไม่มีในเอกสาร ให้เว้นว่างไว้ (null) ห้ามเดาหรือสร้างข้อมูลที่ไม่มีในเอกสารขึ้นมาเอง
"""


# รองรับ Word/Excel/CSV/TXT (Feedback จริงจากผู้ใช้ 2026-09-04) — Gemini แบบ Vision (ที่
# ใช้กับ PDF/รูปภาพผ่าน files.upload ด้านล่าง) อ่านไฟล์ Office ไม่ได้โดยตรง จึงต้องแตก
# ข้อความออกมาก่อนแล้วส่งเป็น Text Content Part แทนการอัปโหลดไฟล์ดิบ — แยก Path กันตาม
# นามสกุลไฟล์ (ตรวจสอบชนิดไฟล์ที่ Endpoint /documents/upload ให้แล้วชั้นหนึ่ง)
_TEXT_EXTRACT_SUFFIXES = {".docx", ".xlsx", ".csv", ".txt"}


def _extract_docx_text(file_path: str) -> str:
    document = docx.Document(file_path)
    lines = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _cell_to_str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _extract_xlsx_text(file_path: str) -> str:
    workbook = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
    lines: list[str] = []
    for sheet in workbook.worksheets:
        lines.append(f"--- Sheet: {sheet.title} ---")
        for row in sheet.iter_rows(values_only=True):
            cells = [_cell_to_str(v) for v in row]
            if any(cells):
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _extract_csv_text(file_path: str) -> str:
    lines: list[str] = []
    with open(file_path, newline="", encoding="utf-8-sig", errors="replace") as f:
        for row in csv.reader(f):
            if any(cell.strip() for cell in row):
                lines.append(" | ".join(row))
    return "\n".join(lines)


def _extract_txt_text(file_path: str) -> str:
    return Path(file_path).read_text(encoding="utf-8-sig", errors="replace")


def _extract_text_content(file_path: str, suffix: str) -> str:
    """อ่านข้อความจากไฟล์ Office/CSV/TXT ตามนามสกุล — ให้ Exception หลุดออกไปตรงๆ ถ้า
    ไฟล์เสีย/อ่านไม่ได้ (ผู้เรียกจะห่อเป็น GeminiExtractionError เอง ไม่ Retry เพราะ
    อ่านซ้ำก็ผิดเหมือนเดิมแน่ๆ ไม่ใช่ Error ชั่วคราวแบบเรียก Gemini API)"""
    if suffix == ".docx":
        return _extract_docx_text(file_path)
    if suffix == ".xlsx":
        return _extract_xlsx_text(file_path)
    if suffix == ".csv":
        return _extract_csv_text(file_path)
    if suffix == ".txt":
        return _extract_txt_text(file_path)
    raise ValueError(f"ไม่รองรับการแตกข้อความจากไฟล์นามสกุล {suffix}")


# Auto-Retry (เพิ่ม 2026-09-03 หลังเจอจริงตอน UAT Walkthrough): Gemini ตอบ
# "503 UNAVAILABLE... currently experiencing high demand" เป็นระยะ — เป็น Error ชั่วคราว
# ฝั่ง Google เอง ไม่ใช่ปัญหาโค้ดเรา จึง Retry อัตโนมัติแบบ Exponential Backoff ให้ก่อน
# ค่อย Fail จริง (ลดการรบกวนผู้ใช้ให้กด Upload ซ้ำเอง)
_RETRYABLE_HTTP_STATUS_CODES = {429, 500, 502, 503, 504}
_MAX_ATTEMPTS = 3
_BACKOFF_BASE_SECONDS = 2.0  # Exponential: รอ 2 วิ แล้ว 4 วิ ระหว่างครั้งที่ Retry


def _is_retryable(exc: Exception) -> bool:
    """แยก Error ชั่วคราว (ควร Retry) ออกจาก Error ถาวร (Retry ไปก็ได้ผลเดิม เสียเวลา
    ผู้ใช้รอเปล่าๆ) — Retryable: Rate Limit (429), Server Error ฝั่ง Google (5xx เช่น
    503 High Demand ที่เจอจริง), Network Timeout/Connection Error — ไม่ Retryable:
    Client Error อื่น เช่น 404 (พิมพ์ชื่อ Model ผิด/Model เลิกให้บริการแล้ว — ดู
    Bug จริงที่เจอวันเดียวกัน 2026-09-03), 400 (Request ผิด), 401/403 (Auth ผิด)
    """
    if isinstance(exc, genai_errors.APIError):
        return exc.code in _RETRYABLE_HTTP_STATUS_CODES
    _retryable_network_errors = (
        httpx.TimeoutException
        | httpx.ConnectError
        | httpx.ReadError
        | ConnectionError
        | TimeoutError
    )
    if isinstance(exc, _retryable_network_errors):
        return True
    return False


class GeminiExtractionError(Exception):
    """เกิดข้อผิดพลาดระหว่างเรียก Gemini API หรือแปลงผลลัพธ์เป็น JSON ตาม Schema"""


@dataclass(frozen=True)
class TokenUsage:
    """จำนวน Token จริงที่ Gemini ตอบกลับมาจาก Call ล่าสุด (2026-09-04 — สำหรับหน้า
    "ค่าใช้จ่าย AI") output_tokens รวม Thinking Token ด้วย (Billed ในอัตราเดียวกับ
    Output ปกติ ตาม Pricing ของ Google)"""

    prompt_tokens: int
    output_tokens: int
    total_tokens: int


class GeminiExtractionService:
    def __init__(
        self,
        client: genai.Client | None = None,
        model: str | None = None,
        max_attempts: int = _MAX_ATTEMPTS,
        backoff_base_seconds: float = _BACKOFF_BASE_SECONDS,
        sleep_fn: Callable[[float], None] = time.sleep,
    ) -> None:
        if client is None:
            if not settings.gemini_api_key:
                raise GeminiExtractionError("ยังไม่ได้ตั้งค่า GEMINI_API_KEY ใน .env")
            client = genai.Client(api_key=settings.gemini_api_key)
        self._client = client
        self._model = model or settings.gemini_model
        self._max_attempts = max_attempts
        self._backoff_base_seconds = backoff_base_seconds
        # sleep_fn Inject ได้เพื่อให้ Unit Test ไม่ต้องรอจริงตอน Retry (ดู
        # tests/test_gemini_extraction.py)
        self._sleep = sleep_fn
        # Token Usage ของ Call ล่าสุด (2026-09-04 — สำหรับหน้า "ค่าใช้จ่าย AI") ผู้เรียก
        # (เช่น documents.py) อ่านค่านี้ต่อจาก extract() ทันทีเพื่อไปบันทึก Log ค่าใช้จ่าย
        # — เป็น None ถ้ายังไม่เคยเรียกสำเร็จเลย หรือ Fail ก่อนได้ Response กลับมา
        self.last_usage: TokenUsage | None = None

    @property
    def model(self) -> str:
        return self._model

    def extract(self, file_path: str) -> ExtractionResult:
        self.last_usage = None
        # Word/Excel/CSV/TXT (2026-09-04): แตกข้อความออกมาก่อนนอก Retry Loop เพราะการ
        # อ่านไฟล์เสีย/Parse ไม่ได้เป็น Error ถาวร Retry ไปก็ได้ผลเดิม (ต่างจาก Error
        # จาก Gemini API ที่ชั่วคราวได้) — PDF/รูปภาพยังใช้ Vision ผ่าน files.upload
        # เหมือนเดิมทุกประการ (text_content เป็น None)
        suffix = Path(file_path).suffix.lower()
        text_content: str | None = None
        if suffix in _TEXT_EXTRACT_SUFFIXES:
            try:
                text_content = _extract_text_content(file_path, suffix)
            except Exception as exc:  # noqa: BLE001
                raise GeminiExtractionError(f"อ่านเนื้อหาไฟล์ไม่สำเร็จ: {exc}") from exc
            if not text_content.strip():
                raise GeminiExtractionError("ไม่พบข้อความในไฟล์ (ไฟล์อาจว่างเปล่า)")

        response = None
        last_exc: Exception | None = None

        for attempt in range(1, self._max_attempts + 1):
            try:
                if text_content is not None:
                    contents = [_EXTRACTION_PROMPT, text_content]
                else:
                    uploaded_file = self._client.files.upload(file=file_path)
                    contents = [_EXTRACTION_PROMPT, uploaded_file]
                response = self._client.models.generate_content(
                    model=self._model,
                    contents=contents,
                    config=types.GenerateContentConfig(
                        response_mime_type="application/json",
                        response_json_schema=ExtractionResult.model_json_schema(),
                    ),
                )
                # เก็บ Token Usage ทันทีที่ได้ Response กลับมา (ก่อนเช็ค response.text ว่าง/
                # แปลง JSON ผ่านหรือไม่) เพราะ Google อาจ Bill Token ไปแล้วถึงแม้ Response
                # จะว่างเปล่าหรือแปลงเป็น JSON ตาม Schema ไม่ได้ก็ตาม — ให้ตัวเลขค่าใช้จ่าย
                # ใกล้เคียงยอด Bill จริงที่สุด getattr กันไว้เผื่อ Fake Client ใน Test เก่า
                # ไม่มี usage_metadata (ไม่ Error แค่ last_usage ยังเป็น None)
                usage_metadata = getattr(response, "usage_metadata", None)
                if usage_metadata is not None:
                    prompt_tokens = usage_metadata.prompt_token_count or 0
                    output_tokens = (usage_metadata.candidates_token_count or 0) + (
                        usage_metadata.thoughts_token_count or 0
                    )
                    total_tokens = usage_metadata.total_token_count or (
                        prompt_tokens + output_tokens
                    )
                    self.last_usage = TokenUsage(
                        prompt_tokens=prompt_tokens,
                        output_tokens=output_tokens,
                        total_tokens=total_tokens,
                    )
                break
            except Exception as exc:  # noqa: BLE001 - ห่อ Exception ทุกชนิดจาก SDK ภายนอกเป็น Error ของเราเอง
                last_exc = exc
                if _is_retryable(exc) and attempt < self._max_attempts:
                    self._sleep(self._backoff_base_seconds * (2 ** (attempt - 1)))
                    continue
                raise GeminiExtractionError(f"เรียก Gemini API ไม่สำเร็จ: {exc}") from exc

        if response is None:  # กันไว้เผื่อ Logic พลาด (ไม่ควรเกิดขึ้นจริง)
            raise GeminiExtractionError(f"เรียก Gemini API ไม่สำเร็จ: {last_exc}")

        if not response.text:
            raise GeminiExtractionError("Gemini ไม่ส่งข้อมูลกลับมา (Response ว่างเปล่า)")

        try:
            return ExtractionResult.model_validate_json(response.text)
        except Exception as exc:  # noqa: BLE001
            raise GeminiExtractionError(
                f"แปลงผลลัพธ์จาก Gemini เป็น JSON ตาม Schema ไม่สำเร็จ: {exc}"
            ) from exc
