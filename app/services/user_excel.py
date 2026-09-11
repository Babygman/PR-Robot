"""User Management — Excel Download/Upload (User Management Redesign, 2026-09-11)

Pattern เดียวกับ app/services/budget_excel.py ทุกประการ (Parse + Validate + Upsert
แบบ Partial-success ในไฟล์เดียว) — Key จับคู่ตอน Re-upload คือ `email` (Unique จริงใน
ตาราง users อยู่แล้ว)

Format: name | email | password (จำเป็นเฉพาะแถวสร้าง User ใหม่ — แถว Update ไม่ใช้
คอลัมน์นี้เลย ไม่ทับรหัสผ่านเดิม ต้องใช้ปุ่ม "Reset Password" แยกต่างหากเท่านั้นถ้าต้องการ
เปลี่ยนรหัสผ่าน — กันบัค Excel เผลอมีคอลัมน์รหัสผ่านของ User อื่นหลุดไปทับโดยไม่ตั้งใจ)
| department | division | position | is_admin | is_fa | can_view_approvals |
can_view_pr | can_view_ar | can_view_all_pr | can_view_all_ar | is_active

Boolean Column รับได้ทั้ง TRUE/FALSE, 1/0, Y/N (ไม่สนตัวพิมพ์ใหญ่เล็ก) ว่างเปล่า = False
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field

import openpyxl

from app.core.security import hash_password
from app.models import User

_BOOL_COLUMNS = [
    "is_admin",
    "is_fa",
    "can_view_approvals",
    "can_view_pr",
    "can_view_ar",
    "can_view_all_pr",
    "can_view_all_ar",
    "is_active",
]
_TEXT_COLUMNS = ["name", "email", "department", "division", "position"]
_ALL_COLUMNS = ["name", "email", "password", *_TEXT_COLUMNS[2:], *_BOOL_COLUMNS]
# หมายเหตุ: _TEXT_COLUMNS[2:] คือ department/division/position (ตัด name/email ที่ซ้ำออก)
_REQUIRED_COLUMNS = ["name", "email"]


@dataclass
class _RowOutcome:
    row_no: int
    ok: bool
    message: str
    raw_data: dict = field(default_factory=dict)


def _parse_bool(value, default: bool = False) -> bool:
    if value is None or value == "":
        return default
    if isinstance(value, bool):
        return value
    text = str(value).strip().lower()
    return text in {"true", "1", "y", "yes"}


def _read_rows(file_bytes: bytes) -> list[dict]:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    ws = wb.worksheets[0]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return []

    header_map: dict[int, str] = {}
    for idx, cell in enumerate(header_row):
        if cell is None:
            continue
        key = str(cell).strip().lower()
        if key in _ALL_COLUMNS:
            header_map[idx] = key

    rows: list[dict] = []
    for raw in rows_iter:
        if raw is None or all(v is None for v in raw):
            continue
        row = {col: (raw[idx] if idx < len(raw) else None) for idx, col in header_map.items()}
        if not row:
            continue
        rows.append(row)
    return rows


def parse_and_upsert(db, *, filename: str, file_bytes: bytes) -> tuple[list[_RowOutcome], int, int]:
    """คืนค่า (outcomes, success_rows, error_rows) — ไม่มี Batch Log Table แยกต่างหาก
    เหมือน Budget (ไม่ได้ร้องขอ) Caller เป็นคน commit() เอง (Pattern เดียวกับ Route อื่น)"""
    try:
        rows = _read_rows(file_bytes)
    except Exception as exc:  # noqa: BLE001 — ไฟล์เสีย/ไม่ใช่ Excel จริง
        return [_RowOutcome(row_no=0, ok=False, message=f"อ่านไฟล์ไม่สำเร็จ: {exc}")], 0, 1

    outcomes: list[_RowOutcome] = []
    seen_emails: set[str] = set()

    for i, row in enumerate(rows, start=2):  # แถว 1 = Header
        missing = [c for c in _REQUIRED_COLUMNS if not row.get(c)]
        if missing:
            outcomes.append(
                _RowOutcome(row_no=i, ok=False, message=f"ขาดข้อมูลจำเป็น: {', '.join(missing)}")
            )
            continue

        email = str(row["email"]).strip().lower()
        name = str(row["name"]).strip()
        if email in seen_emails:
            outcomes.append(
                _RowOutcome(row_no=i, ok=False, message=f'email "{email}" ซ้ำกับแถวอื่นในไฟล์เดียวกัน')
            )
            continue
        seen_emails.add(email)

        field_values = {
            "department": (str(row["department"]).strip() if row.get("department") else None),
            "division": (str(row["division"]).strip() if row.get("division") else None),
            "position": (str(row["position"]).strip() if row.get("position") else None),
        }
        for col in _BOOL_COLUMNS:
            field_values[col] = _parse_bool(row.get(col))

        existing = db.query(User).filter(User.email == email).first()
        if existing is not None:
            existing.name = name
            for key, value in field_values.items():
                setattr(existing, key, value)
            outcomes.append(
                _RowOutcome(
                    row_no=i, ok=True, message=f"อัปเดตแถวเดิม (email={email}, id={existing.id})"
                )
            )
        else:
            password = str(row["password"]).strip() if row.get("password") else ""
            if len(password) < 8:
                outcomes.append(
                    _RowOutcome(
                        row_no=i,
                        ok=False,
                        message=f'User ใหม่ "{email}" ต้องระบุ password อย่างน้อย 8 ตัวอักษรในไฟล์',
                    )
                )
                continue
            new_user = User(
                name=name,
                email=email,
                password_hash=hash_password(password),
                **field_values,
            )
            db.add(new_user)
            outcomes.append(_RowOutcome(row_no=i, ok=True, message=f"เพิ่ม User ใหม่ (email={email})"))

    success_rows = sum(1 for o in outcomes if o.ok)
    error_rows = sum(1 for o in outcomes if not o.ok)
    return outcomes, success_rows, error_rows


def build_export_workbook(users: list[User]) -> bytes:
    """Export User ทั้งหมดเป็น Excel — คอลัมน์ตรงกับ Format Upload เป๊ะ (ยกเว้น password
    ที่ไม่ Export ออกมาด้วยเหตุผลด้านความปลอดภัย — ไฟล์นี้เอาไป Re-upload กลับเข้าระบบได้
    เลยสำหรับแก้ไข User เดิมที่มีอยู่แล้ว แต่สร้าง User ใหม่จากไฟล์นี้ตรงๆ ไม่ได้ ต้องเติม
    คอลัมน์ password เองก่อน)"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Users"
    export_columns = [c for c in _ALL_COLUMNS if c != "password"]
    ws.append(export_columns)
    for u in users:
        ws.append(
            [
                u.name,
                u.email,
                u.department or "",
                u.division or "",
                u.position or "",
                u.is_admin,
                u.is_fa,
                u.can_view_approvals,
                u.can_view_pr,
                u.can_view_ar,
                u.can_view_all_pr,
                u.can_view_all_ar,
                u.is_active,
            ]
        )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
